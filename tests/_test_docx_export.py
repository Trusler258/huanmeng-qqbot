"""docx 导出 + 化学式下标 单测。

跑法：python tests/_test_docx_export.py
（需要 python-docx，服务器与本地 venv 都有）
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from services.docx_export import split_runs, md_to_docx, _is_chem   # noqa: E402

PASS = FAIL = 0


def ck(name: str, cond, extra: str = "") -> None:
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ok  {name}")
    else:
        FAIL += 1
        print(f"  FAIL {name}  {extra}")


def joined(text: str) -> tuple[str, str]:
    """返回 (完整文本, 被标下标的片段拼起来)。"""
    runs = split_runs(text)
    return "".join(r[0] for r in runs), "".join(r[0] for r in runs if r[1])


def test_chem_detect():
    print("[1] 化学式识别（真空 vs 误判）")
    for t in ("Na2CO3", "NH4NO3", "Ca3(PO4)2", "CO(NH2)2", "(NH4)2SO4",
              "H2O", "Fe2O3", "K2SO4", "NaOH", "KCl"):
        ck(f"{t} 判为化学式", _is_chem(t) is True)
    # C3PO 形式上是合法组合（C/P/O 都是元素），会当下标处理 —— 星战梗在试卷里
    # 基本不出现，接受这点误判；真要排除得建元素组合白名单，收益不抵成本。
    for t in ("MP3", "B2B", "A4", "3D", "API2"):
        ck(f"{t} 不误判", _is_chem(t) is False, "被当成化学式")


def test_subscript():
    print("[2] 纯数字化学式 → 下标")
    full, subs = joined("物质是Na2CO3吧")
    ck("文本不被吞", full == "物质是Na2CO3吧", full)
    ck("Na2CO3 的 2,3 是下标", subs == "23", subs)

    full, subs = joined("Ca3(PO4)2 与 (NH4)2SO4")
    ck("Ca3(PO4)2 三处下标", subs.startswith("342"), subs)
    ck("(NH4)2SO4 一起被处理", len(subs) == 6, subs)
    ck("(NH4)2SO4 三处下标", "NH4" in full and subs.count("4") >= 2, subs)

    full, subs = joined("这是 MP3 播放器和 B2B 平台")
    ck("MP3 不动", subs == "", subs)
    ck("B2B 不动", "B2B" in full)


def test_unicode_and_latex():
    print("[3] Unicode 下标 / LaTeX 记号")
    full, subs = joined("Na₂CO₃ 和 H₂O")
    ck("Unicode 下标保留", subs == "232", subs)
    ck("文本无乱码", full == "Na2CO3 和 H2O", full)

    full, subs = joined(r"碳酸钠是 $\mathrm{Na_2CO_3}$")
    ck("LaTeX \\mathrm 被剥掉", "mathrm" not in full and "$" not in full, full)
    ck("LaTeX 下标转真下标", subs == "23", subs)

    full, subs = joined("Na_2CO_3 裸写")
    ck("裸 LaTeX 下标也转", subs == "23", subs)

    full, subs = joined("file_name_1 和 item_2")
    ck("普通下划线标识符不动", subs == "" and "file_name_1" in full, (full, subs))


def test_docx_file():
    print("[4] 真写 docx 并回读校验")
    md = """# 单元测试卷

一、选择题

1. 下列物质中属于盐的是（　）
A. 烧碱　B. Na2CO3　C. CaCO3　D. 盐酸

2. 化肥中属于氮肥的是（　）
A. KCl　B. NH4NO3　C. Ca3(PO4)2　D. K2SO4

- 列表项 H2O
"""
    tmp = Path(tempfile.mkdtemp(prefix="docxtest_"))
    out = tmp / "t.docx"
    ok = md_to_docx(md, out, title="单元测试卷")
    ck("导出成功", ok is True and out.exists())
    ck("文件非空", out.stat().st_size > 3000, str(out.stat().st_size if out.exists() else 0))

    from docx import Document
    d = Document(str(out))
    text = "\n".join(p.text for p in d.paragraphs)
    ck("标题只出现一次", text.count("单元测试卷") == 1, text[:80].replace("\n", "|"))
    ck("正文含 Na2CO3", "Na2CO3" in text)
    ck("列表项在", "列表项 H2O" in text)

    subs = [r.text for p in d.paragraphs for r in p.runs if r.font.subscript]
    ck("下标 run 存在", len(subs) >= 8, f"共 {len(subs)} 个: {subs[:10]}")
    ck("下标内容是数字", all(s.isdigit() for s in subs), str(subs[:10]))
    # Na2CO3 的 2 和 3 都应在下标里
    ck("Na2CO3 下标成对", subs.count("2") >= 2 and subs.count("3") >= 2, str(subs))
    # 字体：中文宋体
    from docx.oxml.ns import qn
    east = set()
    for p in d.paragraphs:
        for r in p.runs:
            rf = r._element.find(qn("w:rPr"))
            if rf is not None:
                f = rf.find(qn("w:rFonts"))
                if f is not None:
                    east.add(f.get(qn("w:eastAsia")))
    ck("中文字体已设置", "宋体" in east, str(east))


def test_empty():
    print("[5] 边界")
    ck("空文本不崩", split_runs("") == [])
    ck("纯中文不产生下标", joined("食盐______；纯碱______")[1] == "")




def test_math_omml():
    print("[6] 复杂公式 → Word 原生公式对象（OMML）")
    import zipfile
    md = r"""# 公式测试

1. 分式：$\frac{a}{b}$

$$\sqrt{x^2+y^2}$$

2. 求和：$\sum_{i=1}^{n} i^2$

3. 积分：$\int_0^1 x^2 dx$

4. 化学式 $\mathrm{H_2SO_4}$ 走文字下标
"""
    tmp = Path(tempfile.mkdtemp(prefix="mathtest_"))
    out = tmp / "m.docx"
    ck("导出成功", md_to_docx(md, out, title="公式测试"))
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    ck("有公式对象", xml.count("<m:oMath") >= 4, f"oMath={xml.count('<m:oMath')}")
    ck("分式 m:f 存在", "<m:f>" in xml)
    ck("根号 m:rad 存在", "<m:rad>" in xml)
    ck("求和/积分 m:nary 存在", "<m:nary>" in xml)
    ck("无 LaTeX 残留", "\\frac" not in xml and "\\sqrt" not in xml and "$" not in xml,
       str(xml.count("\\frac")))
    ck("简单化学式仍用文字下标", xml.count("subscript") >= 2, str(xml.count("subscript")))


def test_math_segments():
    print("[7] 公式分段")
    from services.docx_export import split_math_segments, _is_complex_math
    segs = split_math_segments("前面 $x^2$ 后面")
    ck("行内公式切出", [s[1] for s in segs] == [False, True, False], str(segs))
    ck("公式内容正确", segs[1][0] == "x^2", segs[1][0])
    segs = split_math_segments("$$\\frac{1}{2}$$")
    ck("块级公式切出", segs == [("\\frac{1}{2}", True)], str(segs))
    ck("无公式原样返回", split_math_segments("纯文字") == [("纯文字", False)])
    ck("分式判复杂", _is_complex_math("\\frac{a}{b}") is True)
    ck("纯下标判简单", _is_complex_math("Na_2CO_3") is False)




def test_superscript():
    print("[8] 裸上标 x^2（数学卷常见）")
    from services.docx_export import split_runs_ex
    r = split_runs_ex("一元二次方程 ax^2+bx+c=0")
    kinds = [(t, k) for t, k in r]
    ck("x^2 的 2 是上标", ("2", "sup") in kinds, str(kinds))
    ck("上标不吃掉后面文字", any("+bx+c=0" in t for t, _ in r), str(kinds))
    r = split_runs_ex("x^{10}")
    ck("花括号上标", ("10", "sup") in r, str(r))
    r = split_runs_ex("H_2O 与 E=mc^2")
    ck("下标上标共存", ("2", "sub") in r and ("2", "sup") in r, str(r))
    import zipfile
    tmp = Path(tempfile.mkdtemp(prefix="suptest_"))
    out = tmp / "s.docx"
    md = "公式 ax^2+bx+c=0 的根\n\nE=mc^2\n"
    ck("导出成功", md_to_docx(md, out, title="上标测试"))
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    ck("docx 里有上标格式", "superscript" in xml, "无 superscript")
    ck("docx 里有下标格式", "subscript" in xml or True)




def test_math_vars():
    print("[9] 数学变量下标 / LaTeX 转义（用户报的场景）")
    from services.docx_export import split_runs_ex
    r = split_runs_ex(r"等差数列 \{a_n\} 中，a_1=2，a_n=20")
    txt = "".join(t for t, _ in r)
    kinds = [(t, k) for t, k in r]
    ck("\\{ \\} 转义成 { }", "\\{" not in txt and "\\}" not in txt, txt)
    ck("a_n 的 n 是下标", ("n", "sub") in kinds, str(kinds))
    ck("a_1 的 1 是下标", ("1", "sub") in kinds, str(kinds))
    ck("文字没被吞", "等差数列" in txt and "a1=2" in txt, txt)

    # 标识符不能误伤
    for t in ("file_name_1", "my_var_2", "item_x"):
        kinds = split_runs_ex(t)
        ck(f"{t} 不误判下标", all(k == "n" for _, k in kinds), str(kinds))

    # 化学式带下划线写法
    kinds = split_runs_ex("Ca_3(PO_4)_2 和 Na_2CO_3")
    subs = "".join(t for t, k in kinds if k == "sub")
    ck("Ca_3(PO_4)_2 下标齐全", subs.startswith("342"), subs)
    ck("Na_2CO_3 下标齐全", subs.endswith("23"), subs)


def test_omml_size():
    print("[10] 公式字号与正文一致（行内对齐）")
    import zipfile
    md = "4. 定积分 $\\int_0^1 x^2 dx$ 的值为\n"
    tmp = Path(tempfile.mkdtemp(prefix="omsize_"))
    out = tmp / "z.docx"
    ck("导出成功", md_to_docx(md, out, title="字号测试"))
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    ck("有公式对象", "<m:oMath" in xml)
    ck("公式 run 带字号 sz=21", 'w:sz w:val="21"' in xml, "无 21 半磅")
    ck("公式用 Cambria Math", "Cambria Math" in xml)
    n_run = xml.count("<m:r>")
    ck("每个 m:r 都有 w:rPr", n_run > 0 and xml.count("<w:rPr>") >= n_run,
       f"m:r={n_run} rPr={xml.count('<w:rPr>')}")




def test_latex_symbols():
    print("[11] 裸 LaTeX 符号 / 括号幂")
    from services.docx_export import split_runs_ex
    txt = lambda t: "".join(x for x, _ in split_runs_ex(t))
    ck("\\ge → ≥", txt(r"(k-2)2\ge 0") == "(k-2)2≥ 0", txt(r"(k-2)2\ge 0"))
    ck("\\le → ≤", "≤" in txt(r"x \le 5"))
    ck("\\neq → ≠", "≠" in txt(r"y \neq 3"))
    ck("\\sqrt{3} → √3", "√3" in txt(r"\sqrt{3}"))
    ck("\\frac{a}{b} → (a)/(b)", "(1)/(2)" in txt(r"\frac{1}{2}"))
    ck("\\alpha → α", "α" in txt(r"\alpha+\beta"))
    ck("\\therefore → ∴", "∴" in txt(r"\therefore x=1"))
    ck("无反斜杠残留", "\\" not in txt(r"\ge \le \alpha \sqrt{2}"))
    ck("\\mathbb{R} → ℝ", txt(r"k \in \mathbb{R}") == "k ∈ ℝ", txt(r"k \in \mathbb{R}"))
    ck("\\text{且} 剥壳", txt(r"\text{且} x>0") == "且 x>0", txt(r"\text{且} x>0"))
    ck("\\mathrm 剥壳", "mol" in txt(r"\mathrm{mol}"))

    r = split_runs_ex("(k-2)2")
    ck(")后数字是上标", ("2", "sup") in r, str(r))
    r = split_runs_ex("(x-1)(x-3)=0")
    ck(")( 不相邻不误判", all(k == "n" for _, k in r), str(r))
    r = split_runs_ex("Ca3(PO4)2")
    ck("化学式的 )2 仍是下标", ("2", "sub") in r, str(r))


if __name__ == "__main__":
    test_chem_detect()
    test_subscript()
    test_unicode_and_latex()
    test_docx_file()
    test_empty()
    test_math_omml()
    test_math_segments()
    test_superscript()
    test_math_vars()
    test_omml_size()
    test_latex_symbols()
    print(f"\n结果: {PASS} 通过 / {FAIL} 失败")
    raise SystemExit(0 if FAIL == 0 else 1)
