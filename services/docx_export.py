"""Markdown → Word(.docx) 导出 + 化学式真下标。

为什么要有：
  2026-10-21 群里要"人教版九年级化学第四单元试卷，导出为word"，bot 给了 .md，
  用户拿 Word 打开满屏 `Na2CO3`（下角数字没下标）。md 里没法表达真下标，
  只能靠 Word 的 run.font.subscript。

化学式识别（避免把普通数字当化学式）：
  1. 先用「元素符号 + 可选数字」或「(元素+数字)+数字」拼成 2 段以上的 token
  2. 校验 token 里出现的元素符号**全在白名单**，且不同元素 ≥ 2 种
     → Na2CO3 / NH4NO3 / Ca3(PO4)2 / CO(NH2)2 命中
     → MP3（M 不是元素）、B2B（只有 1 种元素）、A4（A 不是元素）不命中
"""

from __future__ import annotations

import re
from pathlib import Path

from core.logger import get_logger

logger = get_logger(__name__)

# 常见元素符号（覆盖中学化学 + 常见工业元素）
_ELEMENTS = {
    "H", "He", "Li", "Be", "B", "C", "N", "O", "F", "Ne", "Na", "Mg", "Al", "Si",
    "P", "S", "Cl", "Ar", "K", "Ca", "Sc", "Ti", "V", "Cr", "Mn", "Fe", "Co", "Ni",
    "Cu", "Zn", "Ga", "Ge", "As", "Se", "Br", "Kr", "Rb", "Sr", "Y", "Zr", "Nb",
    "Mo", "Tc", "Ru", "Rh", "Pd", "Ag", "Cd", "In", "Sn", "Sb", "Te", "I", "Xe",
    "Cs", "Ba", "La", "Ce", "W", "Pt", "Au", "Hg", "Tl", "Pb", "Bi", "Rn", "Ra",
}

# ⚠️ 括号里可能有两段以上元素："(PO4)" 的 P 后跟大写 O，`[A-Z][a-z]?\d*` 单个
#    片段吃不掉（`[a-z]?` 不吃大写）→ 括号内必须允许重复片段，否则 Ca3(PO4)2
#    只会匹配到 "PO4"，前面的 Ca3 和后面的 2 全丢（实测踩到）。
_ATOM = r"[A-Z][a-z]?\d*"
_GROUP = r"(?:" + _ATOM + r"|\((?:" + _ATOM + r")+\)\d*)"
_CHEM_RE = re.compile(_GROUP + r"{2,}")

# Unicode 下标字符（模型常直接写 Na₂CO₃）
_SUB_CHARS = "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎"
_SUB_MAP = {c: str(i) for i, c in enumerate("₀₁₂₃₄₅₆₇₈₉")}
_REV_SUB = {str(i): c for i, c in enumerate("₀₁₂₃₄₅₆₇₈₉")}

# ★ 用户要求「Word 里 LaTeX 就行」：模型常写 $\mathrm{Na_2CO_3}$ / $H_2O$。
#   Word 本身不认 LaTeX（要手动 Alt+= 粘贴再转换），所以这里直接把它**渲染成真下标**，
#   打开就是对的，用户不用再手工转公式。同时去掉 $ 和 \mathrm{} 这些残留符号。
_LATEX_INLINE_RE = re.compile(r"\$([^$]{1,200})\$")
_LATEX_WRAP_RE = re.compile(r"\\(?:mathrm|text|ce|mathbf|rm|mathit)\s*\{([^{}]*)\}")


def _latex_inline_to_plain(m) -> str:
    """$...$ 里的 LaTeX → 带 Unicode 下标的纯文本。"""
    body = m.group(1)
    body = _LATEX_WRAP_RE.sub(r"\1", body)
    # _{23} → ₂₃
    body = re.sub(r"_\{(\d+)\}",
                  lambda x: "".join(_REV_SUB.get(c, c) for c in x.group(1)), body)
    # _2 → ₂
    body = re.sub(r"_(\d)", lambda x: _REV_SUB.get(x.group(1), x.group(1)), body)
    # ^{2} / ^2 → 普通数字（Word 上标这里不做，避免和下标规则打架）
    body = re.sub(r"\^\{?(\d)\}?", r"\1", body)
    # 剩下的 \ 命令直接去掉反斜杠，保留可读文字
    body = re.sub(r"\\[a-zA-Z]+", "", body)
    return body.strip()


def _is_chem(token: str) -> bool:
    """token 是否是真化学式（白名单元素 + 至少 2 种元素）。"""
    if len(token) < 2:
        return False
    syms = re.findall(r"[A-Z][a-z]?", token)
    if len(set(syms)) < 2:
        return False
    return all(s in _ELEMENTS for s in syms)


# 裸写的上标（数学公式里 x^2 很常见，模型常不带 $）
_SUP_RE = re.compile(r"\^\{([^{}]{1,20})\}|\^([A-Za-z0-9\u0370-\u03ff])")
# 裸写的下标：a_1 / a_n / a_{n+1}（数学变量，不是化学式）
_SUB_MARK_RE = re.compile(r"_\{([^{}]{1,12})\}|_([A-Za-z0-9\u0370-\u03ff])")
# 右括号后紧跟的数字是幂：(k-2)2 → (k-2)²（模型常漏写 ^，实测答案里出现过）
_PAREN_SUP_RE = re.compile(r"\)(\d{1,2})(?!\d)")
# "_" 前面必须是「独立单字母」（a_1）或「元素符号」（Ca_3 / PO_4），或右括号。
# ⚠️ 不能写成 [A-Za-z]{1,2}$ —— 会把 my_var_2 的 _v 当成下标（实测踩到）。
_VAR_BEFORE_SUB_RE = re.compile(
    r"(?<![A-Za-z])[A-Za-z]$|(?<![A-Za-z])[A-Z][A-Za-z]?$|\)$")
# LaTeX 转义字符（模型常写 \{a_n\}，不处理就会原样显示反斜杠）
_LATEX_ESC = {
    "\\{": "{", "\\}": "}", "\\%": "%", "\\&": "&",
    "\\#": "#", "\\$": "$", "\\_": "_",
}


# 裸文本里的 LaTeX 数学符号（模型有时不包 $，直接写 \ge、\times）
# 不转的话 Word 里就是字面 "\ge"（2026-10-21 实测：(k-2)2\ge 0 漏出来）
_LATEX_SYMBOLS = {
    "\\geqslant": "≥", "\\leqslant": "≤", "\\geq": "≥", "\\leq": "≤",
    "\\ge": "≥", "\\le": "≤",
    "\\neq": "≠", "\\ne": "≠", "\\equiv": "≡", "\\approx": "≈",
    "\\pm": "±", "\\mp": "∓", "\\times": "×", "\\div": "÷",
    "\\cdot": "·", "\\infty": "∞", "\\propto": "∝",
    "\\therefore": "∴", "\\because": "∵", "\\angle": "∠",
    "\\triangle": "△", "\\perp": "⊥", "\\parallel": "∥",
    "\\notin": "∉", "\\subset": "⊂", "\\subseteq": "⊆",
    "\\supset": "⊃", "\\in": "∈", "\\cup": "∪", "\\cap": "∩",
    "\\emptyset": "∅", "\\varnothing": "∅", "\\forall": "∀", "\\exists": "∃",
    "\\sum": "∑", "\\prod": "∏", "\\int": "∫", "\\partial": "∂",
    "\\rightarrow": "→", "\\leftarrow": "←", "\\to": "→",
    "\\Rightarrow": "⇒", "\\Leftrightarrow": "⇔", "\\iff": "⇔",
    "\\ldots": "…", "\\cdots": "⋯", "\\dots": "…", "\\circ": "°",
    "\\alpha": "α", "\\beta": "β", "\\gamma": "γ", "\\delta": "δ",
    "\\theta": "θ", "\\lambda": "λ", "\\mu": "μ", "\\pi": "π",
    "\\sigma": "σ", "\\omega": "ω", "\\varphi": "φ", "\\rho": "ρ",
    "\\quad": " ", "\\qquad": "  ",
}


def _strip_latex_escapes(text: str) -> str:
    # 数集：\mathbb{R} → ℝ（模型写 k∈\mathbb{R} 很常见）
    if "\\mathbb" in text:
        _bb = {"R": "ℝ", "N": "ℕ", "Z": "ℤ", "Q": "ℚ", "C": "ℂ"}
        text = re.sub(r"\\mathbb\{([RNZQC])\}",
                      lambda m: _bb.get(m.group(1), m.group(1)), text)
    # 带参数的先处理：\sqrt{x} → √x，\frac{a}{b} → (a)/(b)
    if "\\sqrt" in text:
        text = re.sub(r"\\sqrt\{([^{}]*)\}", r"√\1", text)
        text = re.sub(r"\\sqrt\s*([A-Za-z0-9])", r"√\1", text)
    if "\\frac" in text:
        text = re.sub(r"\\[dt]?frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", text)
    for k in sorted(_LATEX_SYMBOLS, key=len, reverse=True):
        if k in text:
            text = text.replace(k, _LATEX_SYMBOLS[k])
    for k, v in _LATEX_ESC.items():
        if k in text:
            text = text.replace(k, v)
    # 通用带参命令：\text{且} / \mathrm{mol} → 只留花括号里的内容
    text = re.sub(r"\\[a-zA-Z]+\{([^{}]*)\}", r"\1", text)
    # 兜底：没认出来的命令去掉反斜杠，至少不留 "\xxx" 字面
    text = re.sub(r"\\([a-zA-Z]+)", r"\1", text)
    return text


def split_runs_ex(text: str) -> list[tuple[str, str]]:
    """把一段文本切成 [(片段, 类型)]，类型：n 普通 / sub 下标 / sup 上标。

    处理顺序：LaTeX 内联公式降级 → 裸 LaTeX 下标 → 化学式数字 + Unicode 下标 + 裸上标。
    """
    if not text:
        return []
    if "$" in text:
        text = _LATEX_INLINE_RE.sub(_latex_inline_to_plain, text)
    if "\\" in text:
        text = _strip_latex_escapes(text)
    out: list[tuple[str, str]] = []
    pos = 0

    def push(seg: str, kind: str) -> None:
        if not seg:
            return
        if out and out[-1][1] == kind:
            out[-1] = (out[-1][0] + seg, kind)
        else:
            out.append((seg, kind))

    def emit_chars(seg: str) -> None:
        """字符流：Unicode 下标字符 → sub。"""
        buf = []
        for ch in seg:
            if ch in _SUB_CHARS:
                if buf:
                    push("".join(buf), "n")
                    buf = []
                push(_SUB_MAP.get(ch, ""), "sub")
            else:
                buf.append(ch)
        if buf:
            push("".join(buf), "n")

    def emit_sub(seg: str) -> None:
        """段内挑裸下标：a_1 / a_n / a_{n+1} → sub run。

        化学式数字另有 _CHEM_RE 处理，这里只管「变量_下标」写法。
        """
        if not seg:
            return
        if "_" not in seg:
            emit_chars(seg)
            return
        p3 = 0
        for m in _SUB_MARK_RE.finditer(seg):
            if not _VAR_BEFORE_SUB_RE.search(seg[:m.start()]):
                continue          # file_name_1 这类标识符不动
            if m.start() > p3:
                emit_chars(seg[p3:m.start()])
            push(m.group(1) or m.group(2), "sub")
            p3 = m.end()
        if p3 < len(seg):
            emit_chars(seg[p3:])

    def emit_paren_sup(seg: str) -> None:
        """右括号后的数字当上标：(k-2)2 → (k-2)²。"""
        if not seg or ")" not in seg:
            emit_sub(seg)
            return
        p4 = 0
        for m in _PAREN_SUP_RE.finditer(seg):
            if m.start() > p4:
                emit_sub(seg[p4:m.start()])
            push(")", "n")
            push(m.group(1), "sup")
            p4 = m.end()
        if p4 < len(seg):
            emit_sub(seg[p4:])

    def emit_plain(seg: str) -> None:
        """普通文本：先挑裸上标（x^2），再挑 )数字，再挑裸下标（a_1），剩下走字符流。"""
        if not seg:
            return
        if "^" not in seg:
            emit_paren_sup(seg)
            return
        p2 = 0
        for m in _SUP_RE.finditer(seg):
            if m.start() > p2:
                emit_paren_sup(seg[p2:m.start()])
            push(m.group(1) or m.group(2), "sup")
            p2 = m.end()
        if p2 < len(seg):
            emit_paren_sup(seg[p2:])

    for m in _CHEM_RE.finditer(text):
        tok = m.group(0)
        if not _is_chem(tok):
            continue
        emit_plain(text[pos:m.start()])
        for part in re.finditer(r"(\d+)|([^\d]+)", tok):
            if part.group(1):
                push(part.group(1), "sub")
            else:
                push(part.group(2), "n")
        pos = m.end()
        # ★ 紧跟化学式的下标写法（CO_3 / PO_4）：化学式吃掉 "CO" 后，
        #   若不在这里接手，"_3" 落到 emit_plain 时已看不到前面的元素符号 → 漏转
        _tail = re.match(r"_\{?([A-Za-z0-9]{1,3})\}?", text[pos:])
        if _tail:
            push(_tail.group(1), "sub")
            pos += _tail.end()
    emit_plain(text[pos:])
    return out


def split_runs(text: str) -> list[tuple[str, bool]]:
    """兼容旧接口：[(片段, 是否下标)]（上标在这里按普通文本返回）。"""
    return [(seg, kind == "sub") for seg, kind in split_runs_ex(text)]


# ── Markdown → docx ──────────────────────────────────────

_FONT_CN = "宋体"
_FONT_EN = "Times New Roman"


# ══════════════════════════════════════════════════════════
# 复杂 LaTeX 公式 → Word 原生公式对象（OMML）
# ══════════════════════════════════════════════════════════
# 为什么不能只用下标：下标只能表达 Na2CO3 这种原子式。分式、根号、求和、积分、
# 矩阵在 Word 里必须是 OMML（<m:oMath>）对象，纯文字下标做不出来。
# 链路：LaTeX --latex2mathml--> MathML --mathml2omml--> OMML，插进 run。
# 两个包没装/转换失败 → 自动退回"降级成下标文字"，不至于整篇失败。

_MATH_INLINE_RE = re.compile(r"\$\$([^$]+?)\$\$|\$([^$\n]+?)\$|\\\[(.+?)\\\]|\\\((.+?)\\\)",
                             re.S)
# 出现这些命令 = 简单下标搞不定，必须走 OMML
_COMPLEX_MATH_RE = re.compile(
    r"\\(?:frac|dfrac|tfrac|sqrt|sum|prod|int|iint|oint|lim|begin|end|matrix|"
    r"binom|overline|underline|vec|hat|bar|partial|nabla|alpha|beta|gamma|delta|"
    r"theta|lambda|mu|sigma|pi|omega|times|cdot|div|pm|leq|geq|neq|approx|"
    r"infty|rightarrow|leftarrow|Rightarrow|to|quad|left|right)",
    re.IGNORECASE,
)


def _omml_available() -> bool:
    try:
        import latex2mathml.converter  # noqa: F401
        import mathml2omml  # noqa: F401
        return True
    except Exception:
        return False


def _latex_to_omml(latex: str) -> str:
    """LaTeX → OMML 字符串（失败抛异常）。"""
    import latex2mathml.converter as _l2m
    import mathml2omml as _m2o
    return _m2o.convert(_l2m.convert(latex))


# OMML 里每个公式 run 的开口（后面可能紧跟 <m:rPr>）
_OMML_RUN_RE = re.compile(r"<m:r>(<m:rPr>(?:(?!</m:rPr>).)*</m:rPr>)?")


def _omml_set_size(om: str, size_pt: float) -> str:
    """给每个 m:r 补 w:rPr（Cambria Math + 与正文同字号）。

    ⚠️ 不补的话公式沿用 Word 默认字号（11pt），而正文是 10.5pt：
    公式比文字大一号，行内看起来忽高忽低、跟文字不在一条基线上（用户实测反馈）。
    """
    hp = max(2, int(round(float(size_pt) * 2)))     # w:sz 单位是半磅
    rpr = (
        '<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/>'
        f'<w:sz w:val="{hp}"/><w:szCs w:val="{hp}"/></w:rPr>'
    )
    return _OMML_RUN_RE.sub(lambda m: "<m:r>" + (m.group(1) or "") + rpr, om)


def _append_omml(paragraph, latex: str, size_pt: float = 10.5) -> bool:
    """把公式作为 Word 原生公式对象插进段落。成功 True。"""
    try:
        from docx.oxml import parse_xml
        om = _latex_to_omml(latex)
        if not om or "<m:oMath" not in om:
            return False
        om = _omml_set_size(om, size_pt)
        om = om.replace(
            "<m:oMath",
            '<m:oMath xmlns:m="http://schemas.openxmlformats.org/'
            'officeDocument/2006/math"'
            ' xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"',
            1,
        )
        paragraph._p.append(parse_xml(om))
        return True
    except Exception as e:
        logger.warning("公式转 OMML 失败（改回文本下标）: %s | %s", latex, e)
        return False


def split_math_segments(text: str) -> list[tuple[str, bool]]:
    """按 $...$ / $$...$$ / \\[..\\] / \\(..\\) 切成 [(片段, 是否公式)]。"""
    if not text or ("$" not in text and "\\[" not in text and "\\(" not in text):
        return [(text, False)] if text else []
    out: list[tuple[str, bool]] = []
    pos = 0
    for m in _MATH_INLINE_RE.finditer(text):
        body = m.group(1) or m.group(2) or m.group(3) or m.group(4) or ""
        if not body.strip():
            continue
        if m.start() > pos:
            out.append((text[pos:m.start()], False))
        out.append((body.strip(), True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False))
    return out


def _is_complex_math(latex: str) -> bool:
    """是否必须用 OMML（简单下标交给文字下标更自然）。"""
    return bool(_COMPLEX_MATH_RE.search(latex or ""))


def _write_runs(paragraph, text: str, size_pt: float = 10.5, bold: bool = False) -> None:
    """把一行文本写进段落：公式走 OMML，其余走文字 + 化学式下标。"""
    omml_ok = _omml_available()
    for seg, is_math in split_math_segments(text):
        if is_math:
            if omml_ok and _is_complex_math(seg) and _append_omml(paragraph, seg, size_pt):
                continue
            # 简单公式（Na_2CO_3 这种）或 OMML 不可用 → 降级成下标文字
            # 先剥掉 \mathrm{...} / \text{...} 字体壳，否则会变成字面文字
            text_part = _LATEX_WRAP_RE.sub(r"\1", seg)
        else:
            text_part = seg
        for piece, kind in split_runs_ex(text_part):
            r = paragraph.add_run(piece)
            _set_font(r, size_pt=size_pt, bold=bold)
            if kind == "sub":
                r.font.subscript = True
            elif kind == "sup":
                r.font.superscript = True


def _set_font(run, size_pt: float | None = None, bold: bool = False) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt
    run.font.name = _FONT_EN
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), _FONT_CN)
    if size_pt:
        run.font.size = Pt(size_pt)
    run.bold = bold


def _write_para(doc, text: str, style: str | None = None, size_pt: float = 10.5,
                bold: bool = False, indent_first: bool = False) -> None:
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    if indent_first:
        p.paragraph_format.first_line_indent = p.style.font.size or None
        from docx.shared import Pt
        p.paragraph_format.first_line_indent = Pt(size_pt * 2)
    _write_runs(p, text, size_pt=size_pt, bold=bold)
    return p


def md_to_docx(md_text: str, out_path: str | Path, title: str = "") -> bool:
    """把 markdown 文本写成 .docx。成功 True，失败 False（失败原因写日志）。"""
    try:
        from docx import Document
        from docx.shared import Pt
    except Exception as e:  # pragma: no cover
        logger.error("python-docx 不可用，无法导出 word: %s", e)
        return False

    try:
        doc = Document()
        # 默认正文样式
        doc.styles["Normal"].font.name = _FONT_EN
        doc.styles["Normal"].font.size = Pt(10.5)

        # 标题（文件名去掉扩展名），正文里若已有同名首行则跳过重复
        lines = (md_text or "").split("\n")
        head = title.strip()
        if lines and lines[0].strip():
            first = lines[0].strip().lstrip("#").strip()
            if head and first == head:
                lines = lines[1:]
        if head:
            h = doc.add_paragraph()
            h.alignment = 1  # 居中
            r = h.add_run(head)
            _set_font(r, size_pt=16, bold=True)
            h.paragraph_format.space_after = Pt(10)

        for raw in lines:
            line = raw.rstrip()
            s = line.strip()
            if not s:
                continue
            # 分隔线
            if re.fullmatch(r"[-*_]{3,}", s):
                continue
            # 标题
            m = re.match(r"^(#{1,6})\s+(.*)$", s)
            if m:
                level = min(len(m.group(1)), 4)
                txt = m.group(2).strip()
                size = {1: 15, 2: 13.5, 3: 12, 4: 11}[level]
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(4)
                _write_runs(p, txt, size_pt=size, bold=True)
                continue
            # 无序列表
            m = re.match(r"^[-*+]\s+(.*)$", s)
            if m:
                _write_para(doc, m.group(1), style="List Bullet")
                continue
            # 表格行（简单处理：竖线分隔的跳过表头分隔线，其余按制表符分隔成一行）
            if s.startswith("|") and s.endswith("|"):
                if re.fullmatch(r"\|[\s:|-]+\|", s):
                    continue
                cells = [c.strip() for c in s.strip("|").split("|")]
                _write_para(doc, "　".join(cells))
                continue
            # ★ 块级公式（整行只有一个 $$...$$ 或 \\[...\\]）→ 单独居中段落
            _blk = re.fullmatch(r"\$\$(.+?)\$\$|\\\[(.+?)\\\]", s, re.S)
            if _blk:
                body = (_blk.group(1) or _blk.group(2) or "").strip()
                p = doc.add_paragraph()
                p.alignment = 1
                if not (_omml_available() and _append_omml(p, body, 11)):
                    _write_runs(p, body, size_pt=11)
                continue
            # 普通段落（试卷题号、选项行都走这里，保留原样，不自动编号）
            _write_para(doc, s)

        Path(out_path).parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out_path))
        logger.info("docx 导出成功: %s（%d 段）", Path(out_path).name, len(doc.paragraphs))
        return True
    except Exception as e:
        logger.error("docx 导出失败: %s", e, exc_info=True)
        return False
